"""Service for local file conversion, transformation, and document processing.

Fully local: uses PyMuPDF, Pillow, python-docx, openpyxl, ffmpeg, ImageMagick,
and LibreOffice to execute actual conversions without sending files to external APIs.
"""
from __future__ import annotations

import asyncio
import mimetypes
import os
import re
import shutil
import time
import zipfile
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import UploadFile

from backend.config import paths

# Max upload size: 500MB for media/documents
MAX_CONVERTER_UPLOAD_BYTES = 500 * 1024 * 1024

IMAGE_EXTENSIONS = {
    "png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff", "tif", "ico", "avif", "svg",
}
DOCUMENT_EXTENSIONS = {
    "pdf", "docx", "doc", "xlsx", "xls", "pptx", "ppt", "txt", "md", "csv", "html", "rtf", "odt",
}
VIDEO_EXTENSIONS = {
    "mp4", "webm", "mkv", "mov", "avi", "wmv", "flv", "m4v", "mpeg", "mpg",
}
AUDIO_EXTENSIONS = {
    "mp3", "wav", "m4a", "aac", "ogg", "flac", "opus", "wma", "aiff",
}
ARCHIVE_EXTENSIONS = {
    "zip", "tar", "gz", "tgz", "bz2", "xz", "7z", "tar.gz", "tar.bz2", "tar.xz",
}


def _probe_media(source: Path) -> dict[str, Any]:
    """Inspect video/audio stream properties using ffprobe."""
    if not _check_binary("ffprobe"):
        return {"has_audio": True, "has_video": True}

    import json
    import subprocess

    cmd = [
        "ffprobe", "-v", "quiet", "-print_format", "json",
        "-show_streams", "-show_format", str(source),
    ]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=10)
        data = json.loads(res.stdout)
        streams = data.get("streams", [])
        has_audio = any(s.get("codec_type") == "audio" for s in streams)
        has_video = any(s.get("codec_type") == "video" for s in streams)
        return {
            "has_audio": has_audio,
            "has_video": has_video,
            "streams": streams,
        }
    except Exception:
        return {"has_audio": True, "has_video": True}


def _get_category(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext == "pdf":
        return "pdf"
    if ext in DOCUMENT_EXTENSIONS:
        return "document"
    if ext in VIDEO_EXTENSIONS:
        return "video"
    if ext in AUDIO_EXTENSIONS:
        return "audio"
    if ext in ARCHIVE_EXTENSIONS or ext.endswith(("tar.gz", "tar.bz2", "tar.xz", "tgz", "zip")):
        return "archive"
    return "other"


def _check_binary(name: str) -> bool:
    return shutil.which(name) is not None


class ConverterService:
    def __init__(self) -> None:
        self.base_dir = paths().home / "converter"
        self.uploads_dir = self.base_dir / "uploads"
        self.outputs_dir = self.base_dir / "outputs"
        self.uploads_dir.mkdir(parents=True, exist_ok=True)
        self.outputs_dir.mkdir(parents=True, exist_ok=True)

    def get_capabilities(self) -> dict[str, Any]:
        """Inspect local environment and return available tools, formats, and operations."""
        has_ffmpeg = _check_binary("ffmpeg")
        has_magick = _check_binary("magick") or _check_binary("convert")
        has_soffice = _check_binary("soffice") or _check_binary("libreoffice")

        try:
            import fitz  # noqa: F401
            has_pymupdf = True
        except ImportError:
            has_pymupdf = False

        try:
            import PIL  # noqa: F401
            has_pillow = True
        except ImportError:
            has_pillow = False

        try:
            import docx  # noqa: F401
            has_docx = True
        except ImportError:
            has_docx = False

        try:
            import openpyxl  # noqa: F401
            has_openpyxl = True
        except ImportError:
            has_openpyxl = False

        try:
            from rapidocr_onnxruntime import RapidOCR  # noqa: F401
            has_ocr = True
        except ImportError:
            has_ocr = False

        engines = {
            "ffmpeg": has_ffmpeg,
            "imagemagick": has_magick,
            "libreoffice": has_soffice,
            "pymupdf": has_pymupdf,
            "pillow": has_pillow,
            "docx": has_docx,
            "openpyxl": has_openpyxl,
            "ocr": has_ocr,
        }

        # Determine supported target formats per category
        image_targets = ["png", "jpg", "webp", "gif", "bmp", "tiff", "ico", "pdf", "txt"]
        if has_magick:
            image_targets.append("avif")

        video_targets = []
        if has_ffmpeg:
            video_targets = ["mp4", "webm", "mkv", "mov", "gif", "mp3", "wav", "aac"]

        audio_targets = []
        if has_ffmpeg:
            audio_targets = ["mp3", "wav", "aac", "ogg", "flac", "m4a", "opus"]

        doc_targets = ["pdf", "txt", "md"]
        if has_soffice:
            doc_targets.extend(["docx", "xlsx", "html"])

        return {
            "engines": engines,
            "categories": {
                "image": {
                    "label": "Images",
                    "formats": sorted(list(IMAGE_EXTENSIONS)),
                    "targets": image_targets,
                    "operations": [
                        "convert", "resize", "compress", "rotate", "grayscale", "extract_text",
                    ],
                },
                "pdf": {
                    "label": "PDF Documents",
                    "formats": ["pdf"],
                    "targets": ["images_zip", "images_png", "txt", "docx", "pdf_compressed"],
                    "operations": [
                        "pdf_to_images", "pdf_merge", "pdf_extract_pages",
                        "pdf_extract_text", "pdf_compress", "convert",
                    ],
                },
                "document": {
                    "label": "Office & Documents",
                    "formats": sorted(list(DOCUMENT_EXTENSIONS - {"pdf"})),
                    "targets": doc_targets,
                    "operations": ["convert", "extract_text"],
                },
                "video": {
                    "label": "Video Files",
                    "formats": sorted(list(VIDEO_EXTENSIONS)),
                    "targets": video_targets,
                    "operations": ["convert", "video_to_audio", "video_to_gif", "video_compress"],
                },
                "audio": {
                    "label": "Audio Files",
                    "formats": sorted(list(AUDIO_EXTENSIONS)),
                    "targets": audio_targets,
                    "operations": ["convert"],
                },
                "archive": {
                    "label": "Archives & Compressed",
                    "formats": sorted(list(ARCHIVE_EXTENSIONS)),
                    "targets": ["zip", "tar.gz", "tar"],
                    "operations": ["convert"],
                },
            },
            "max_file_size_bytes": MAX_CONVERTER_UPLOAD_BYTES,
        }

    def generate_thumbnail(self, source: Path, thumb_path: Path) -> bool:
        """Extract a single representative video frame to JPEG using FFmpeg."""
        if not _check_binary("ffmpeg"):
            return False
        import subprocess

        for t in ["00:00:01", "00:00:00"]:
            cmd = [
                "ffmpeg", "-y", "-ss", t, "-i", str(source),
                "-vframes", "1", "-vf", "scale=480:-1", str(thumb_path),
            ]
            res = subprocess.run(
                cmd,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=10,
            )
            if res.returncode == 0 and thumb_path.exists() and thumb_path.stat().st_size > 0:
                return True
        return False

    def ensure_thumbnail(self, file_id: str) -> Path | None:
        """Find or generate a cached thumbnail for an uploaded media file."""
        upload_path = self.get_upload_path(file_id)
        if not upload_path or not upload_path.is_file():
            return None
        folder = self.uploads_dir / file_id
        thumb_path = folder / "thumbnail.jpg"
        if thumb_path.is_file():
            return thumb_path

        cat = _get_category(upload_path.suffix)
        if cat == "video":
            if self.generate_thumbnail(upload_path, thumb_path):
                return thumb_path
        elif cat == "image":
            try:
                from PIL import Image
                with Image.open(upload_path) as img:
                    img.thumbnail((480, 480))
                    img.convert("RGB").save(thumb_path, "JPEG")
                return thumb_path
            except Exception:
                return None
        return None

    async def save_upload(self, file: UploadFile) -> dict[str, Any]:
        """Save an uploaded file to disk and return file metadata."""
        file_id = uuid4().hex[:14]
        original_name = Path(file.filename or "upload").name
        safe_name = re.sub(r"[^\w\s\-.]", "_", original_name).strip() or f"file_{file_id}"
        ext = safe_name.split(".")[-1].lower() if "." in safe_name else ""

        target_dir = self.uploads_dir / file_id
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / safe_name

        size = 0
        with target_path.open("wb") as out:
            while chunk := await file.read(1 << 20):  # 1MB chunks
                size += len(chunk)
                if size > MAX_CONVERTER_UPLOAD_BYTES:
                    out.close()
                    shutil.rmtree(target_dir, ignore_errors=True)
                    max_mb = MAX_CONVERTER_UPLOAD_BYTES // (1024 * 1024)
                    raise ValueError(f"File exceeds maximum allowed size of {max_mb}MB")
                out.write(chunk)

        category = _get_category(ext)
        guessed = mimetypes.guess_type(str(target_path))[0]
        mime_type = file.content_type or guessed or "application/octet-stream"

        has_audio = True
        has_thumbnail = False
        if category == "video":
            probe = _probe_media(target_path)
            has_audio = probe.get("has_audio", True)
            thumb_path = target_dir / "thumbnail.jpg"
            if self.generate_thumbnail(target_path, thumb_path):
                has_thumbnail = True
        elif category == "image":
            has_thumbnail = True

        return {
            "id": file_id,
            "file_id": file_id,
            "filename": safe_name,
            "size": size,
            "ext": ext,
            "category": category,
            "mime_type": mime_type,
            "has_audio": has_audio,
            "has_thumbnail": has_thumbnail,
            "thumbnail_url": f"/api/converter/thumbnail/{file_id}" if has_thumbnail else None,
            "uploaded_at": int(time.time()),
        }

    def get_upload_path(self, file_id: str) -> Path | None:
        """Find the local path for a given file_id."""
        folder = self.uploads_dir / file_id
        if not folder.is_dir():
            return None
        files = list(folder.glob("*"))
        if not files:
            return None
        return files[0]

    def get_output_path(self, job_id: str) -> Path | None:
        """Find the processed file for a given job_id."""
        folder = self.outputs_dir / job_id
        if not folder.is_dir():
            return None
        files = list(folder.glob("*"))
        if not files:
            return None
        return files[0]

    async def process(
        self,
        file_id: str | None = None,
        file_ids: list[str] | None = None,
        operation: str = "convert",
        target_format: str | None = None,
        options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Execute processing or conversion operation."""
        options = options or {}
        job_id = uuid4().hex[:14]
        job_dir = self.outputs_dir / job_id
        job_dir.mkdir(parents=True, exist_ok=True)

        start_time = time.time()

        # Multi-file operations
        if operation == "pdf_merge":
            if not file_ids or len(file_ids) < 2:
                raise ValueError("Merging PDFs requires at least 2 files.")
            result_file = await self._merge_pdfs(file_ids, job_dir, options)
        elif operation == "images_to_pdf":
            if not file_ids or len(file_ids) < 1:
                raise ValueError("Images to PDF requires at least 1 image.")
            result_file = await self._images_to_pdf(file_ids, job_dir, options)
        else:
            # Single file operations
            if not file_id:
                raise ValueError("file_id is required for this operation.")
            source_path = self.get_upload_path(file_id)
            if not source_path or not source_path.is_file():
                raise FileNotFoundError(f"Uploaded file '{file_id}' not found.")

            source_ext = source_path.suffix.lstrip(".").lower()
            source_category = _get_category(source_ext)

            if operation == "convert":
                if not target_format:
                    raise ValueError("Target format must be specified.")
                target_clean = target_format.lower().lstrip(".")
                result_file = await self._convert_file(
                    source_path, target_clean, job_dir, options,
                )
            elif operation == "resize":
                result_file = await self._resize_image(source_path, job_dir, options)
            elif operation == "compress":
                if source_category == "video":
                    result_file = await self._video_compress(source_path, job_dir, options)
                else:
                    result_file = await self._compress_image(source_path, job_dir, options)
            elif operation == "rotate":
                result_file = await self._rotate_image(source_path, job_dir, options)
            elif operation == "grayscale":
                result_file = await self._grayscale_image(source_path, job_dir, options)
            elif operation == "pdf_to_images":
                result_file = await self._pdf_to_images(source_path, job_dir, options)
            elif operation == "pdf_extract_pages":
                result_file = await self._pdf_extract_pages(source_path, job_dir, options)
            elif operation in {"pdf_extract_text", "extract_text"}:
                result_file = await self._extract_text(source_path, job_dir, options)
            elif operation == "pdf_compress":
                result_file = await self._pdf_compress(source_path, job_dir, options)
            elif operation == "video_to_audio":
                audio_format = options.get("format", target_format or "mp3").lower().lstrip(".")
                result_file = await self._video_to_audio(
                    source_path, audio_format, job_dir, options,
                )
            elif operation == "video_to_gif":
                result_file = await self._video_to_gif(source_path, job_dir, options)
            elif operation == "video_compress":
                result_file = await self._video_compress(source_path, job_dir, options)
            else:
                raise ValueError(f"Unknown operation: {operation}")

        elapsed_ms = int((time.time() - start_time) * 1000)
        output_size = result_file.stat().st_size
        output_ext = result_file.suffix.lstrip(".").lower()
        output_mime = mimetypes.guess_type(str(result_file))[0] or "application/octet-stream"

        preview_text = None
        if output_ext in {"txt", "md", "csv", "json", "html"} and output_size < 100 * 1024:
            try:
                preview_text = result_file.read_text(encoding="utf-8", errors="replace")
            except Exception:
                preview_text = None

        return {
            "job_id": job_id,
            "status": "completed",
            "operation": operation,
            "filename": result_file.name,
            "size": output_size,
            "ext": output_ext,
            "mime_type": output_mime,
            "elapsed_ms": elapsed_ms,
            "download_url": f"/api/converter/download/{job_id}",
            "preview_url": f"/api/converter/preview/{job_id}",
            "preview_text": preview_text,
        }

    # =========================================================================
    # Operation Handlers
    # =========================================================================

    async def _convert_file(
        self, source: Path, target: str, job_dir: Path, options: dict[str, Any],
    ) -> Path:
        source_ext = source.suffix.lstrip(".").lower()
        source_cat = _get_category(source_ext)
        target_cat = _get_category(target)

        dest_name = f"{source.stem}.{target}"
        dest_path = job_dir / dest_name

        # 0. Converting anything to txt: use dedicated extractor
        if target == "txt":
            return await self._extract_text(source, job_dir, options)

        # 1. Archives (zip, tar, tar.gz)
        if source_cat == "archive" or target_cat == "archive" or target in {"zip", "tar.gz", "tar"}:
            return await asyncio.to_thread(self._archive_convert, source, dest_path, target)

        # 2. Image to Image or Image to PDF via Pillow
        if source_cat == "image" and (target_cat == "image" or target == "pdf"):
            return await asyncio.to_thread(self._pillow_convert, source, dest_path, target, options)

        # 3. Audio/Video conversions via FFmpeg
        if (source_cat in {"video", "audio"}) and (target_cat in {"video", "audio", "image"}):
            if target == "gif" and source_cat == "video":
                return await self._video_to_gif(source, job_dir, options)
            if target_cat == "audio" and source_cat == "video":
                return await self._video_to_audio(source, target, job_dir, options)
            return await self._ffmpeg_convert(source, dest_path, target, options)

        # 4. Document / Office conversion via LibreOffice headless
        if source_cat in {"document", "pdf"} or target_cat in {"document", "pdf"}:
            if target == "pdf" and source_cat == "image":
                return await asyncio.to_thread(
                    self._pillow_convert, source, dest_path, "pdf", options,
                )
            return await self._libreoffice_convert(source, dest_path, target)

        # Fallback to ImageMagick if available
        if _check_binary("magick"):
            proc = await asyncio.create_subprocess_exec(
                "magick", str(source), str(dest_path),
                stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
            )
            _, err = await proc.communicate()
            if proc.returncode == 0 and dest_path.exists():
                return dest_path
            raise RuntimeError(f"ImageMagick failed: {err.decode(errors='replace')}")

        raise ValueError(f"Cannot convert '{source.name}' to format '{target}'.")

    def _pillow_convert(
        self, source: Path, dest: Path, target: str, options: dict[str, Any],
    ) -> Path:
        from PIL import Image

        fmt_map = {
            "jpg": "JPEG",
            "jpeg": "JPEG",
            "png": "PNG",
            "webp": "WEBP",
            "gif": "GIF",
            "bmp": "BMP",
            "tiff": "TIFF",
            "ico": "ICO",
            "pdf": "PDF",
        }
        pil_fmt = fmt_map.get(target.lower(), target.upper())

        with Image.open(source) as img:
            # Handle RGBA to RGB for JPEG/PDF
            if pil_fmt in {"JPEG", "PDF"} and img.mode in ("RGBA", "LA", "P"):
                background = Image.new("RGB", img.size, (255, 255, 255))
                if img.mode == "P":
                    img = img.convert("RGBA")
                background.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
                save_img = background
            else:
                save_img = img

            save_kwargs: dict[str, Any] = {}
            if pil_fmt == "JPEG":
                quality = int(options.get("quality", 90))
                save_kwargs = {"quality": quality, "optimize": True}
            elif pil_fmt == "WEBP":
                quality = int(options.get("quality", 85))
                save_kwargs = {"quality": quality, "method": 6}
            elif pil_fmt == "PNG":
                save_kwargs = {"optimize": True}
            elif pil_fmt == "ICO":
                sizes = [(16, 16), (32, 32), (48, 48), (64, 64)]
                save_kwargs = {"sizes": sizes}

            save_img.save(dest, format=pil_fmt, **save_kwargs)

        return dest

    async def _resize_image(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        dest_path = job_dir / source.name

        def _do_resize():
            from PIL import Image
            with Image.open(source) as img:
                orig_w, orig_h = img.size

                scale_pct = options.get("percentage")
                width = options.get("width")
                height = options.get("height")
                maintain_aspect = options.get("maintain_aspect", True)

                if scale_pct:
                    pct = float(scale_pct) / 100.0
                    new_w = max(1, int(orig_w * pct))
                    new_h = max(1, int(orig_h * pct))
                elif width and height:
                    new_w = int(width)
                    new_h = int(height)
                    if maintain_aspect:
                        img.thumbnail((new_w, new_h), Image.Resampling.LANCZOS)
                        img.save(dest_path)
                        return
                elif width:
                    new_w = int(width)
                    new_h = max(1, int(orig_h * (new_w / orig_w)))
                elif height:
                    new_h = int(height)
                    new_w = max(1, int(orig_w * (new_h / orig_h)))
                else:
                    new_w, new_h = orig_w, orig_h

                resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
                resized.save(dest_path)

        await asyncio.to_thread(_do_resize)
        return dest_path

    async def _compress_image(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        dest_path = job_dir / source.name

        def _do_compress():
            from PIL import Image
            quality = int(options.get("quality", 75))
            with Image.open(source) as img:
                fmt = img.format or "JPEG"
                if fmt == "PNG":
                    img.save(dest_path, format="PNG", optimize=True)
                elif fmt in {"JPEG", "JPG"}:
                    img.save(dest_path, format="JPEG", quality=quality, optimize=True)
                elif fmt == "WEBP":
                    img.save(dest_path, format="WEBP", quality=quality, method=6)
                else:
                    img.save(dest_path, format=fmt, quality=quality)

        await asyncio.to_thread(_do_compress)
        return dest_path

    async def _rotate_image(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        dest_path = job_dir / source.name
        angle = int(options.get("angle", 90))

        def _do_rotate():
            from PIL import Image
            with Image.open(source) as img:
                # Pillow rotate is counter-clockwise by default, so invert angle for clockwise UX
                rotated = img.rotate(-angle, expand=True)
                rotated.save(dest_path)

        await asyncio.to_thread(_do_rotate)
        return dest_path

    async def _grayscale_image(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        dest_path = job_dir / source.name

        def _do_grayscale():
            from PIL import Image
            with Image.open(source) as img:
                gray = img.convert("L")
                gray.save(dest_path)

        await asyncio.to_thread(_do_grayscale)
        return dest_path

    async def _pdf_to_images(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        """Render PDF pages as images (PNG or JPG) and return single image or ZIP."""
        import fitz

        img_format = options.get("format", "png").lower().lstrip(".")
        dpi = int(options.get("dpi", 200))
        zoom = dpi / 72.0
        mat = fitz.Matrix(zoom, zoom)

        doc = fitz.open(source)
        page_count = len(doc)
        if page_count == 0:
            doc.close()
            raise ValueError("PDF contains no pages.")

        if page_count == 1:
            page = doc.load_page(0)
            pix = page.get_pixmap(matrix=mat)
            dest_file = job_dir / f"{source.stem}_page_1.{img_format}"
            pix.save(str(dest_file))
            doc.close()
            return dest_file

        # Multiple pages: package into ZIP
        zip_path = job_dir / f"{source.stem}_pages.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for idx in range(page_count):
                page = doc.load_page(idx)
                pix = page.get_pixmap(matrix=mat)
                img_data = pix.tobytes(output=img_format)
                zf.writestr(f"page_{idx + 1:03d}.{img_format}", img_data)

        doc.close()
        return zip_path

    async def _merge_pdfs(
        self, file_ids: list[str], job_dir: Path, options: dict[str, Any],
    ) -> Path:
        """Merge multiple PDF files into one combined PDF document."""
        import fitz

        output_name = options.get("output_name") or "merged_document.pdf"
        if not output_name.endswith(".pdf"):
            output_name += ".pdf"
        dest_path = job_dir / output_name

        merged_doc = fitz.open()
        for fid in file_ids:
            path = self.get_upload_path(fid)
            if not path or not path.is_file():
                continue
            sub_doc = fitz.open(path)
            merged_doc.insert_pdf(sub_doc)
            sub_doc.close()

        merged_doc.save(str(dest_path))
        merged_doc.close()
        return dest_path

    async def _images_to_pdf(
        self, file_ids: list[str], job_dir: Path, options: dict[str, Any],
    ) -> Path:
        """Convert one or more images into a single clean PDF document."""
        import fitz

        output_name = options.get("output_name") or "images_combined.pdf"
        if not output_name.endswith(".pdf"):
            output_name += ".pdf"
        dest_path = job_dir / output_name

        pdf_doc = fitz.open()
        for fid in file_ids:
            path = self.get_upload_path(fid)
            if not path or not path.is_file():
                continue
            img_doc = fitz.open(path)
            rect = img_doc[0].rect
            pdf_bytes = img_doc.convert_to_pdf()
            img_doc.close()

            img_pdf = fitz.open("pdf", pdf_bytes)
            page = pdf_doc.new_page(width=rect.width, height=rect.height)
            page.show_pdf_page(rect, img_pdf, 0)
            img_pdf.close()

        pdf_doc.save(str(dest_path))
        pdf_doc.close()
        return dest_path

    async def _pdf_extract_pages(
        self, source: Path, job_dir: Path, options: dict[str, Any],
    ) -> Path:
        """Extract given page numbers or ranges (e.g. '1, 3-5') into a new PDF."""
        import fitz

        pages_spec = str(options.get("pages", "1")).strip()
        doc = fitz.open(source)
        total = len(doc)

        selected_indices = set()
        for part in pages_spec.split(","):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                start_s, end_s = part.split("-", 1)
                s = max(1, int(start_s.strip()))
                e = min(total, int(end_s.strip()))
                for p in range(s, e + 1):
                    selected_indices.add(p - 1)
            else:
                p = int(part)
                if 1 <= p <= total:
                    selected_indices.add(p - 1)

        if not selected_indices:
            doc.close()
            raise ValueError(f"No valid pages found in range '{pages_spec}'. Total pages: {total}")

        sorted_indices = sorted(list(selected_indices))
        new_doc = fitz.open()
        for idx in sorted_indices:
            new_doc.insert_pdf(doc, from_page=idx, to_page=idx)

        dest_path = job_dir / f"{source.stem}_pages_{pages_spec.replace(' ', '')}.pdf"
        new_doc.save(str(dest_path))
        new_doc.close()
        doc.close()
        return dest_path

    async def _pdf_compress(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        """Compress PDF file size using Ghostscript if available, or PyMuPDF deflating."""
        dest_path = job_dir / f"{source.stem}_compressed.pdf"

        if _check_binary("gs"):
            proc = await asyncio.create_subprocess_exec(
                "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
                "-dPDFSETTINGS=/ebook", "-dNOPAUSE", "-dQUIET", "-dBATCH",
                f"-sOutputFile={dest_path}", str(source),
                stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
            )
            await proc.communicate()
            if proc.returncode == 0 and dest_path.exists() and dest_path.stat().st_size > 0:
                return dest_path

        # PyMuPDF fallback compression
        import fitz
        doc = fitz.open(source)
        doc.save(str(dest_path), garbage=4, deflate=True, clean=True)
        doc.close()
        return dest_path

    def _ocr_image(self, source: Path) -> str:
        """Run OCR on image using RapidOCR."""
        try:
            from rapidocr_onnxruntime import RapidOCR
            engine = RapidOCR()
            result, _ = engine(str(source))
            if not result:
                return "(No text detected in image via OCR)"
            boxes = []
            for item in result:
                if len(item) < 2 or not item[1].strip():
                    continue
                coords, text = item[0], item[1].strip()
                xs, ys = [pt[0] for pt in coords], [pt[1] for pt in coords]
                boxes.append({
                    "y_center": (min(ys) + max(ys)) / 2,
                    "x_min": min(xs),
                    "x_max": max(xs),
                    "height": max(ys) - min(ys),
                    "text": text
                })

            if not boxes:
                return "(No text detected in image via OCR)"

            boxes.sort(key=lambda b: b["y_center"])
            lines, current_line = [], []
            current_y, line_height = None, 0

            for b in boxes:
                if current_y is None:
                    current_line.append(b)
                    current_y, line_height = b["y_center"], b["height"]
                else:
                    if abs(b["y_center"] - current_y) < (line_height * 0.5):
                        current_line.append(b)
                        current_y = sum(x["y_center"] for x in current_line) / len(current_line)
                        line_height = sum(x["height"] for x in current_line) / len(current_line)
                    else:
                        lines.append(current_line)
                        current_line = [b]
                        current_y, line_height = b["y_center"], b["height"]
            if current_line:
                lines.append(current_line)

            total_chars = sum(len(b["text"]) for b in boxes)
            avg_char_width = sum(b["x_max"] - b["x_min"] for b in boxes) / total_chars if total_chars > 0 else 10

            output_lines = []
            for line in lines:
                line.sort(key=lambda b: b["x_min"])
                line_str, current_x = "", 0
                for b in line:
                    spaces = max(0, int((b["x_min"] - current_x) / avg_char_width))
                    if 0 < spaces < 3:
                        line_str += " "
                    elif spaces >= 3:
                        line_str += " " * spaces
                    line_str += b["text"]
                    current_x = b["x_max"]
                output_lines.append(line_str.rstrip())

            return "\n".join(output_lines)
        except ImportError:
            try:
                import fitz
                # Convert the image to a PDF first so get_textpage_ocr() can work
                imgdoc = fitz.open(source)
                pdfbytes = imgdoc.convert_to_pdf()
                imgdoc.close()
                doc = fitz.open("pdf", pdfbytes)
                tp = doc[0].get_textpage_ocr()
                text = tp.extractText()
                doc.close()
                return text.strip() or "(No text detected in image)"
            except Exception as e:
                raise RuntimeError(f"OCR requires rapidocr-onnxruntime, and PyMuPDF fallback failed: {e}")
        except Exception as exc:
            raise RuntimeError(f"OCR extraction error: {exc}")

    def _archive_convert(self, source: Path, dest: Path, target: str) -> Path:
        """Convert between archive formats (zip, tar, tar.gz) using Python standard library."""
        import tarfile
        import tempfile

        temp_extract = Path(tempfile.mkdtemp(prefix="amethyst_archive_"))
        try:
            src_name = source.name.lower()
            def _extract_tar(tf, path):
                if hasattr(tarfile, "data_filter"):
                    tf.extractall(path, filter="data")
                else:
                    tf.extractall(path)

            if src_name.endswith(".zip"):
                with zipfile.ZipFile(source, "r") as zf:
                    zf.extractall(temp_extract)
            elif src_name.endswith((".tar.gz", ".tgz")):
                with tarfile.open(source, "r:gz") as tf:
                    _extract_tar(tf, temp_extract)
            elif src_name.endswith(".tar"):
                with tarfile.open(source, "r:") as tf:
                    _extract_tar(tf, temp_extract)
            elif src_name.endswith((".tar.bz2", ".tbz2")):
                with tarfile.open(source, "r:bz2") as tf:
                    _extract_tar(tf, temp_extract)
            elif src_name.endswith((".tar.xz", ".txz")):
                with tarfile.open(source, "r:xz") as tf:
                    _extract_tar(tf, temp_extract)
            else:
                raise ValueError(f"Unsupported source archive format: {source.suffix}")

            target_clean = target.lower().lstrip(".")
            if target_clean == "zip":
                with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zf:
                    for root, _, files in os.walk(temp_extract):
                        for f in files:
                            p = Path(root) / f
                            rel = p.relative_to(temp_extract)
                            zf.write(p, arcname=str(rel))
            elif target_clean in {"tar.gz", "tgz"}:
                with tarfile.open(dest, "w:gz") as tf:
                    for root, _, files in os.walk(temp_extract):
                        for f in files:
                            p = Path(root) / f
                            rel = p.relative_to(temp_extract)
                            tf.add(p, arcname=str(rel))
            elif target_clean == "tar":
                with tarfile.open(dest, "w:") as tf:
                    for root, _, files in os.walk(temp_extract):
                        for f in files:
                            p = Path(root) / f
                            rel = p.relative_to(temp_extract)
                            tf.add(p, arcname=str(rel))
            else:
                raise ValueError(f"Unsupported target archive format: {target}")

            return dest
        finally:
            shutil.rmtree(temp_extract, ignore_errors=True)

    async def _extract_text(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        """Extract text from documents (PDF, Word, Excel, OCR images, text files)."""
        source_ext = source.suffix.lstrip(".").lower()
        dest_path = job_dir / f"{source.stem}_extracted.txt"

        extracted = ""
        if source_ext in IMAGE_EXTENSIONS:
            extracted = await asyncio.to_thread(self._ocr_image, source)
        elif source_ext == "pdf":
            import fitz
            doc = fitz.open(source)
            pages = []
            for i, page in enumerate(doc):
                txt = page.get_text().strip()
                if not txt:
                    # Fallback to OCR on page pixmap for scanned PDFs
                    try:
                        import tempfile
                        pix = page.get_pixmap(dpi=150)
                        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                            tmp_path = Path(tmp.name)
                        pix.save(str(tmp_path))
                        ocr_txt = self._ocr_image(tmp_path)
                        tmp_path.unlink(missing_ok=True)
                        if ocr_txt and not ocr_txt.startswith("("):
                            txt = ocr_txt
                    except Exception:
                        pass
                if txt:
                    pages.append(f"--- Page {i + 1} ---\n{txt}")
            doc.close()
            extracted = "\n\n".join(pages)
        elif source_ext == "docx":
            import docx
            doc = docx.Document(source)
            extracted = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif source_ext in {"xlsx", "xls"}:
            import openpyxl
            wb = openpyxl.load_workbook(source, data_only=True)
            sheet_texts = []
            for name in wb.sheetnames:
                sheet = wb[name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell) if cell is not None else "" for cell in row]
                    if any(row_vals):
                        rows.append("\t".join(row_vals))
                if rows:
                    sheet_texts.append(f"=== Sheet: {name} ===\n" + "\n".join(rows))
            wb.close()
            extracted = "\n\n".join(sheet_texts)
        elif source_ext in {"txt", "md", "csv", "json", "html", "xml", "log", "py", "js", "css"}:
            extracted = source.read_text(encoding="utf-8", errors="replace")
        else:
            extracted = f"(Format '{source_ext}' cannot be extracted as text)"

        if not extracted.strip():
            extracted = "(No readable text content found in document)"

        dest_path.write_text(extracted, encoding="utf-8")
        return dest_path

    async def _ffmpeg_convert(
        self, source: Path, dest: Path, target: str, options: dict[str, Any],
    ) -> Path:
        if not _check_binary("ffmpeg"):
            raise RuntimeError(
                "ffmpeg is required for audio and video conversion but is not installed on PATH."
            )

        cmd = ["ffmpeg", "-y", "-i", str(source)]

        # Specific video codecs
        if target == "mp4":
            cmd.extend(["-c:v", "libx264", "-preset", "fast", "-crf", "23", "-c:a", "aac"])
        elif target == "webm":
            cmd.extend(["-c:v", "libvpx-vp9", "-crf", "30", "-b:v", "0", "-c:a", "libopus"])
        elif target == "mp3":
            cmd.extend(["-vn", "-acodec", "libmp3lame", "-q:a", "2"])
        elif target == "wav":
            cmd.extend(["-vn"])
        elif target == "aac":
            cmd.extend(["-vn", "-acodec", "aac", "-b:a", "192k"])

        cmd.append(str(dest))

        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        _, err = await proc.communicate()
        if proc.returncode == 0 and dest.exists():
            return dest
        raise RuntimeError(f"FFmpeg conversion failed: {err.decode(errors='replace')[-500:]}")

    async def _video_to_audio(
        self, source: Path, audio_format: str, job_dir: Path, options: dict[str, Any],
    ) -> Path:
        if not _check_binary("ffmpeg"):
            raise RuntimeError("ffmpeg is not installed on PATH.")

        probe = _probe_media(source)
        if not probe.get("has_audio", True):
            raise ValueError(
                f"Uploaded video '{source.name}' does not contain an audio track to extract."
            )

        dest_path = job_dir / f"{source.stem}.{audio_format}"
        cmd = ["ffmpeg", "-y", "-i", str(source), "-vn", "-map", "0:a:0?"]

        if audio_format == "mp3":
            bitrate = options.get("bitrate", "192k")
            cmd.extend(["-c:a", "libmp3lame", "-b:a", str(bitrate)])
        elif audio_format == "aac":
            cmd.extend(["-c:a", "aac", "-b:a", "192k"])
        elif audio_format == "wav":
            cmd.extend(["-c:a", "pcm_s16le"])
        elif audio_format == "flac":
            cmd.extend(["-c:a", "flac"])
        elif audio_format == "ogg":
            cmd.extend(["-c:a", "libvorbis", "-q:a", "5"])

        cmd.append(str(dest_path))

        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        _, err = await proc.communicate()
        if proc.returncode == 0 and dest_path.exists():
            return dest_path
        raise RuntimeError(f"Audio extraction failed: {err.decode(errors='replace')[-500:]}")

    async def _video_to_gif(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        """High-quality 2-pass palette-based GIF generator using FFmpeg."""
        if not _check_binary("ffmpeg"):
            raise RuntimeError("ffmpeg is not installed on PATH.")

        dest_path = job_dir / f"{source.stem}.gif"
        fps = int(options.get("fps", 15))
        width = int(options.get("width", 480))

        # Using palettegen + paletteuse filter
        filter_str = (
            f"fps={fps},scale={width}:-1:flags=lanczos,split[s0][s1];"
            f"[s0]palettegen[p];[s1][p]paletteuse"
        )
        cmd = ["ffmpeg", "-y", "-i", str(source), "-vf", filter_str, str(dest_path)]

        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        _, err = await proc.communicate()
        if proc.returncode == 0 and dest_path.exists():
            return dest_path
        raise RuntimeError(f"GIF generation failed: {err.decode(errors='replace')[-500:]}")

    async def _video_compress(self, source: Path, job_dir: Path, options: dict[str, Any]) -> Path:
        """Compress video using h264 CRF encoding with audio/subtitle compatibility."""
        if not _check_binary("ffmpeg"):
            raise RuntimeError("ffmpeg is not installed on PATH.")

        dest_path = job_dir / f"{source.stem}_compressed.mp4"
        crf = str(options.get("crf", 28))  # 28 is high compression, 23 is default
        preset = options.get("preset", "fast")

        probe = _probe_media(source)
        has_audio = probe.get("has_audio", True)

        cmd = [
            "ffmpeg", "-y", "-i", str(source),
            "-sn",  # strip subtitles so unsupported subtitles do not break MP4 container
            "-c:v", "libx264", "-crf", crf, "-preset", preset,
        ]

        if has_audio:
            cmd.extend(["-c:a", "aac", "-b:a", "128k"])
        else:
            cmd.append("-an")

        cmd.append(str(dest_path))

        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        _, err = await proc.communicate()
        if proc.returncode == 0 and dest_path.exists():
            return dest_path
        raise RuntimeError(f"Video compression failed: {err.decode(errors='replace')[-500:]}")

    async def _libreoffice_convert(self, source: Path, dest: Path, target: str) -> Path:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise RuntimeError(
                f"Converting '{source.name}' to '{target}' requires LibreOffice, "
                "which is not installed."
            )

        outdir = dest.parent
        profile_dir = self.base_dir / "soffice_profile"
        profile_dir.mkdir(parents=True, exist_ok=True)
        user_install = f"-env:UserInstallation={profile_dir.as_uri()}"

        cmd = [
            soffice, "--headless", "--norestore", user_install,
        ]

        source_ext = source.suffix.lstrip(".").lower()
        target_clean = target.lstrip(".").lower()
        # PDF to Word DOCX requires writer_pdf_import so LibreOffice uses Writer rather than Draw
        if source_ext == "pdf" and target_clean in {"docx", "doc", "odt", "rtf"}:
            cmd.append("--infilter=writer_pdf_import")

        cmd.extend(["--convert-to", target_clean, "--outdir", str(outdir), str(source)])

        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        _, err = await proc.communicate()

        # LibreOffice outputs source.stem + .target in outdir
        expected = outdir / f"{source.stem}.{target_clean}"
        if expected.exists():
            if expected != dest:
                expected.replace(dest)
            return dest

        raise RuntimeError(f"LibreOffice conversion failed: {err.decode(errors='replace')[-500:]}")
